"""Histórico de North Carolina (Pick 3 + Pick 4) y California (Daily 3) a Excel.

Uso: python extract_nc_ca.py north_carolina <carpeta_csv> <salida.xlsx>
     python extract_nc_ca.py california <cache.json> <salida.xlsx>

North Carolina: descarga oficial de la NC Education Lottery, un CSV con todo el histórico:
  https://nclottery.com/pick3-download y https://nclottery.com/pick4-download
Day (15:00) → Tarde y Evening (23:22) → Noche. La Fireball (desde el 18/09/2022) va en dos
columnas aparte que el wizard ignora. Los Double Draw de Pick 3 (segundo sorteo
promocional de algunas noches, marcados "D") no son el sorteo regular: van a Observaciones.
El wizard pide los premios 2 y 3, así que la hoja principal arranca cuando hay Pick 3 y
Pick 4 juntos; los años con solo Pick 3 van a la hoja "Solo Pick 3".

California: calottery.com responde 403 desde Uruguay, así que se usa lotterycorner
(Midday y Evening) cruzado con lottery.net (desde 2010), como Kentucky: si difieren manda
lottery.net. Solo Daily 3: centena y número, sin premios 2 y 3. No tiene bola extra.
"""
import csv
import json
import os
import re
import sys
import time
from collections import Counter
from datetime import date, datetime, timedelta

import requests
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from extract_kentucky import parsear_lc, DIAS

UA = ('Mozilla/5.0 (Windows NT 10.0; Win64; x64) '
      'AppleWebKit/537.36 (KHTML, like Gecko) Chrome/147.0 Safari/537.36')
TURNOS = {'midday': ('afternoon', 'Tarde'), 'evening': ('evening', 'Noche')}
DELAY = 2.0

URL_NC = 'https://nclottery.com/pick%d-download'
URL_CA_LC = 'https://www.lotterycorner.com/ca/daily-3-%s/%d'
URL_CA_LNET = 'https://www.lottery.net/california/daily-3-%s/numbers/%d'
PRIMER_ANIO_CA_LC = 1992
PRIMER_ANIO_CA_LNET = 2010

RE_FILA_LNET = re.compile(r'<tr>\s*<td class="colour noBefore pickArchive">\s*([^<]+?)\s*</td>(.*?)</tr>', re.S)
RE_BOLA_LNET = re.compile(r'class="ball ca-daily-3 number-part-0([1-3])">(\d)<')
MESES = {m: i for i, m in enumerate(
    ['January', 'February', 'March', 'April', 'May', 'June', 'July', 'August',
     'September', 'October', 'November', 'December'], 1)}


def parsear_lnet(html, bolas_n=3):
    out = {}
    for fecha_txt, cuerpo in RE_FILA_LNET.findall(html):
        partes = fecha_txt.replace(',', '').split()
        try:
            d = date(int(partes[3]), MESES[partes[1]], int(partes[2]))
        except (ValueError, IndexError, KeyError):
            continue
        bolas = dict(RE_BOLA_LNET.findall(cuerpo))
        if len(bolas) == bolas_n:
            out[d.isoformat()] = ''.join(bolas[str(i)] for i in range(1, bolas_n + 1))
    return out


def hoja(wb, titulo, encabezado, filas, anchos, texto_cols, primera=False):
    ws = wb.active if primera else wb.create_sheet(titulo)
    ws.title = titulo
    ws.append(encabezado)
    for f in filas:
        ws.append(list(f))
    for c in ws[1]:
        c.font = Font(bold=True)
    for row in ws.iter_rows(min_row=2, min_col=1, max_col=1):
        row[0].number_format = 'yyyy-mm-dd'
    for col in texto_cols:
        for row in ws.iter_rows(min_row=2, min_col=col, max_col=col):
            row[0].number_format = '@'
    for i, w in enumerate(anchos, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = 'A2'
    return ws


def huecos(datos, codigo, hasta, obs):
    """Días sin sorteo entre el primero de la fuente y hoy."""
    d = date.fromisoformat(min(datos))
    while d < hasta:
        if d.isoformat() not in datos:
            obs.append((d, codigo, 'sin resultado en la fuente (se omite el turno)', DIAS[d.weekday()]))
        d += timedelta(days=1)


def resumen(filas, obs):
    print('filas', len(filas), 'desde', filas[0][0], 'hasta', filas[-1][0])
    for codigo in ('afternoon', 'evening'):
        fs = [r[0] for r in filas if r[1] == codigo]
        print(codigo, len(fs), fs[0], fs[-1], 'dup', len(fs) - len(set(fs)))
    print('observaciones', Counter(o[2] for o in obs))
    print('huecos por año', sorted(Counter((o[0].year, o[1]) for o in obs if o[2].startswith('sin')).items()))


# ── North Carolina ────────────────────────────────────────────────

def leer_nc(path, juego, obs):
    """→ {turno: {iso: (número, fireball)}} con el sorteo regular."""
    out = {'midday': {}, 'evening': {}}
    with open(path, encoding='latin-1', newline='') as f:
        for r in csv.DictReader(f):
            try:
                d = datetime.strptime(r['Date'], '%m/%d/%Y').date()
            except (ValueError, TypeError):
                continue
            turno = {'D': 'midday', 'E': 'evening'}.get(r['Day/Eve'])
            num = ''.join(r['Ball %d' % i] for i in range(1, juego + 1))
            if (r.get('DoubleDraw*') or '') == 'D':
                obs.append((d, TURNOS[turno or 'evening'][0], 'Double Draw (sorteo promocional extra, no se usa)',
                            'Pick %d: %s' % (juego, num)))
                continue
            if not turno or not re.fullmatch(r'\d{%d}' % juego, num):
                continue
            if d.isoformat() in out[turno]:
                obs.append((d, TURNOS[turno][0], 'fila repetida en la fuente (queda la primera)', 'Pick %d' % juego))
                continue
            out[turno][d.isoformat()] = (num, (r.get('Fireball') or '').strip())
    return out


def north_carolina(carpeta, salida):
    s = requests.Session()
    s.headers['User-Agent'] = UA
    os.makedirs(carpeta, exist_ok=True)
    for juego in (3, 4):
        r = s.get(URL_NC % juego, timeout=120)
        r.raise_for_status()
        with open(os.path.join(carpeta, 'nc_pick%d.csv' % juego), 'wb') as f:
            f.write(r.content)
    obs = []
    p3 = leer_nc(os.path.join(carpeta, 'nc_pick3.csv'), 3, obs)
    p4 = leer_nc(os.path.join(carpeta, 'nc_pick4.csv'), 4, obs)
    hoy = date.today()

    filas, solo3 = [], []
    for turno, (codigo, nombre) in TURNOS.items():
        a, b = p3[turno], p4[turno]
        inicio4 = min(b)
        for k in sorted(a):
            d = date.fromisoformat(k)
            if k in b:
                (n3, fb3), (n4, fb4) = a[k], b[k]
                filas.append((d, codigo, nombre, n3, n4, int(n3[0]), n3[1:], n4[:2], n4[2:],
                              'nclottery.com', fb3, fb4))
            elif k < inicio4:
                solo3.append((d, codigo, nombre, a[k][0], int(a[k][0][0]), a[k][0][1:], a[k][1]))
            else:
                obs.append((d, codigo, 'sorteo incompleto en la fuente (se omite el turno)', 'falta Pick 4'))
        for k in sorted(set(b) - set(a)):
            obs.append((date.fromisoformat(k), codigo, 'sorteo incompleto en la fuente (se omite el turno)',
                        'falta Pick 3'))
        huecos(a, codigo, hoy, obs)

    filas.sort(key=lambda r: (r[0], r[1] != 'afternoon'))
    solo3.sort(key=lambda r: (r[0], r[1] != 'afternoon'))
    wb = Workbook()
    hoja(wb, 'North Carolina',
         ['Fecha', 'Turno', 'Turno (nombre)', 'Pick 3', 'Pick 4', 'Centena', 'Número',
          'Premio 2', 'Premio 3', 'Fuente', 'Fireball P3', 'Fireball P4'],
         filas, [12, 10, 15, 8, 8, 9, 9, 9, 9, 14, 11, 11], (4, 5, 7, 8, 9, 11, 12), primera=True)
    hoja(wb, 'Solo Pick 3',
         ['Fecha', 'Turno', 'Turno (nombre)', 'Pick 3', 'Centena', 'Número', 'Fireball P3'],
         solo3, [12, 10, 15, 8, 9, 9, 11], (4, 6, 7))
    hoja(wb, 'Observaciones', ['Fecha', 'Turno', 'Observación', 'Detalle'],
         sorted(obs, key=lambda x: (x[0], x[1])), [12, 10, 52, 20], ())
    wb.save(salida)
    resumen(filas, obs)
    print('solo pick 3', len(solo3), Counter(r[1] for r in solo3), solo3[0][0], solo3[-1][0])


# ── California ────────────────────────────────────────────────────

def california(cache_path, salida):
    cache = {}
    if os.path.exists(cache_path):
        with open(cache_path, encoding='utf-8') as f:
            cache = json.load(f)
    hoy = date.today()
    s = requests.Session()
    s.headers['User-Agent'] = UA
    fuentes = (('lc', URL_CA_LC, parsear_lc, PRIMER_ANIO_CA_LC), ('lnet', URL_CA_LNET, parsear_lnet, PRIMER_ANIO_CA_LNET))
    for anio in range(PRIMER_ANIO_CA_LC, hoy.year + 1):
        for fuente, url, parser, primer_anio in fuentes:
            if anio < primer_anio:
                continue
            for turno in TURNOS:
                k = '%s/%s/%d' % (fuente, turno, anio)
                if k in cache and anio < hoy.year:
                    continue
                for intento in range(3):
                    try:
                        r = s.get(url % (turno, anio), timeout=60)
                        cache[k] = {} if r.status_code == 404 else parser(r.raise_for_status() or r.text, 3)
                        print(k, len(cache[k]))
                        break
                    except Exception as exc:
                        print('error', k, exc)
                        time.sleep(5 * (intento + 1))
                time.sleep(DELAY)
        with open(cache_path, 'w', encoding='utf-8') as f:
            json.dump(cache, f)

    filas, obs = [], []
    for turno, (codigo, nombre) in TURNOS.items():
        lc, lnet = {}, {}
        for k, v in cache.items():
            if k.startswith('lc/%s/' % turno):
                lc.update(v)
            elif k.startswith('lnet/%s/' % turno):
                lnet.update(v)
        for k in sorted(set(lc) & set(lnet)):
            if lc[k] != lnet[k]:
                obs.append((date.fromisoformat(k), codigo, 'lotterycorner y lottery.net difieren (queda lottery.net)',
                            '%s / %s' % (lc[k], lnet[k])))
        for k in sorted(set(lnet) - set(lc)):
            obs.append((date.fromisoformat(k), codigo, 'solo en lottery.net (se completa con esa fuente)', 'Daily 3'))
        datos = dict(lc, **lnet)
        if not datos:
            continue
        for k, n in datos.items():
            filas.append((date.fromisoformat(k), codigo, nombre, n, int(n[0]), n[1:],
                          'lottery.net' if k in lnet else 'lotterycorner.com'))
        huecos(datos, codigo, hoy, obs)

    filas.sort(key=lambda r: (r[0], r[1] != 'afternoon'))
    wb = Workbook()
    hoja(wb, 'California', ['Fecha', 'Turno', 'Turno (nombre)', 'Daily 3', 'Centena', 'Número', 'Fuente'],
         filas, [12, 10, 15, 8, 9, 9, 16], (4, 6), primera=True)
    hoja(wb, 'Observaciones', ['Fecha', 'Turno', 'Observación', 'Detalle'],
         sorted(obs, key=lambda x: (x[0], x[1])), [12, 10, 56, 16], ())
    wb.save(salida)
    resumen(filas, obs)


if __name__ == '__main__':
    {'north_carolina': north_carolina, 'california': california}[sys.argv[1]](sys.argv[2], sys.argv[3])
