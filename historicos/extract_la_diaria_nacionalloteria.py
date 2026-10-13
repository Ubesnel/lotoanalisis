"""Histórico completo de La Diaria (Honduras) desde nacionalloteria.com.

Uso: python extract_la_diaria_nacionalloteria.py <cache.json> <salida.xlsx> [excel_api.xlsx]

Una página por día (?del-dia=YYYY-MM-DD), con una tabla de 3 filas:
Horario | Sorteo (nro. correlativo hasta ~2020, después "Día/Tarde/Noche") |
"NN + d" (número + Más 1; antes de ~2018 solo "NN") | Signo.
Primer día con datos: 31/05/2015. Horarios: 10am/2pm/9pm al principio,
11am/3pm/9pm después; se mapean por franja a morning/afternoon/evening.

El caché JSON guarda cada día ya bajado, así que se puede cortar y retomar.
Si se pasa el Excel de la API (extract_la_diaria.py), desde su primera fecha
manda la API (es la fuente de los datos de este sitio, que a veces mete 00 de
relleno o se saltea días) y este sitio solo completa los huecos. Las
diferencias quedan en la hoja Observaciones.

La celda de números viene a veces mal escrita ("11 + 09", "78.Tienda + 3",
"08 + JG", "70 +", "97 + ?"): se toman los dos primeros dígitos. Si el Más 1
trae un número de dos cifras ("00 + 67", "36 + 36") la fila es dudosa y se
descarta.
"""
import html
import json
import os
import re
import sys
import time
from datetime import date, datetime, timedelta

import requests
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

URL = 'https://www.nacionalloteria.com/honduras/loto-la-diaria.php'
FIRST_DAY = date(2015, 5, 31)
DELAY = 1.0
NO_DATA = 'n no est'          # "Aún no están disponibles los premios…"

ROW_RE = re.compile(
    r'<td class="res text-center">([^<]*)</td>\s*<td class="res text-center">([^<]*)</td>'
    r'\s*<td class="res text-center">([^<]*)</td>\s*<td class="res text-center">([^<]*)</td>')
TIME_RE = re.compile(r'<time[^>]*datetime="(\d{4}-\d{2}-\d{2})')
NUM_RE = re.compile(r'^(\d{1,2})(?=$|[\s.+])')
MAS1_RE = re.compile(r'\+\s*(\d{1,2})\s*$')


def turno_de(hora):
    """10/11 am → morning, 2/3 pm → afternoon, 9 pm → evening."""
    m = re.match(r'(\d{1,2}):\d{2}\s*([ap])m', hora.strip().lower())
    if not m:
        return None
    h, ap = int(m.group(1)), m.group(2)
    if ap == 'a':
        return 'morning'
    if h in (12, 1, 2, 3, 4):
        return 'afternoon'
    if h in (8, 9, 10):
        return 'evening'
    return None


def parse_day(texto, pedido):
    texto = html.unescape(texto)
    if NO_DATA in texto:
        return {'status': 'sin_datos', 'rows': []}
    fechas = TIME_RE.findall(texto)
    if fechas and fechas[0] != pedido:
        return {'status': 'otra_fecha', 'rows': [], 'detalle': fechas[0]}
    i = texto.find('Resultados del La Diaria')
    tabla = texto[i:texto.find('</table>', i)] if i >= 0 else ''
    rows = []
    for hora, sorteo, nums, signo in ROW_RE.findall(tabla):
        rows.append({'hora': hora.strip(), 'sorteo': sorteo.strip(),
                     'nums': nums.strip(), 'signo': signo.strip()})
    return {'status': 'ok' if rows else 'sin_tabla', 'rows': rows}


def descargar(cache_path, hasta):
    cache = {}
    if os.path.exists(cache_path):
        with open(cache_path, encoding='utf-8') as f:
            cache = json.load(f)
    s = requests.Session()
    s.headers['User-Agent'] = ('Mozilla/5.0 (Windows NT 10.0; Win64; x64) '
                               'AppleWebKit/537.36 (KHTML, like Gecko) Chrome/147.0 Safari/537.36')
    d, nuevos = FIRST_DAY, 0
    while d <= hasta:
        k = d.isoformat()
        if k not in cache or cache[k]['status'] == 'error':
            for intento in range(3):
                try:
                    r = s.get(URL, params={'del-dia': k}, timeout=40)
                    r.raise_for_status()
                    cache[k] = parse_day(r.text, k)
                    break
                except Exception as exc:
                    cache[k] = {'status': 'error', 'rows': [], 'detalle': str(exc)}
                    time.sleep(5 * (intento + 1))
            nuevos += 1
            if nuevos % 50 == 0:
                with open(cache_path, 'w', encoding='utf-8') as f:
                    json.dump(cache, f, ensure_ascii=False)
                print(k, 'descargados', nuevos, flush=True)
            time.sleep(DELAY)
        d += timedelta(days=1)
    with open(cache_path, 'w', encoding='utf-8') as f:
        json.dump(cache, f, ensure_ascii=False)
    return cache


def main():
    cache_path, out_path = sys.argv[1], sys.argv[2]
    api_path = sys.argv[3] if len(sys.argv) > 3 else None
    cache = descargar(cache_path, date.today() - timedelta(days=1))

    orden = {'morning': 0, 'afternoon': 1, 'evening': 2}
    nombre = {'morning': 'Mañana', 'afternoon': 'Tarde', 'evening': 'Noche'}
    data, obs = [], []
    for k in sorted(cache):
        dia = cache[k]
        fecha = date.fromisoformat(k)
        if dia['status'] != 'ok':
            obs.append((fecha, '', dia['status'], dia.get('detalle', '')))
            continue
        vistos = set()
        for r in dia['rows']:
            turno = turno_de(r['hora'])
            m = NUM_RE.match(r['nums'])
            mas = MAS1_RE.search(r['nums'])
            if not turno or not m:
                obs.append((fecha, turno or '', 'fila ilegible', f"{r['hora']} | {r['nums']}"))
                continue
            if mas and int(mas.group(1)) > 9:
                obs.append((fecha, turno, 'fila dudosa (descartada)', f"{r['hora']} | {r['nums']}"))
                continue
            if turno in vistos:
                obs.append((fecha, turno, 'turno repetido en el día', r['hora']))
                continue
            vistos.add(turno)
            data.append({'fecha': fecha, 'turno': turno, 'hora': r['hora'],
                         'sorteo': r['sorteo'] if r['sorteo'].isdigit() else '',
                         'numero': int(m.group(1)),
                         'mas_1': str(int(mas.group(1))) if mas else '',
                         'signo': r['signo'], 'fuente': 'nacionalloteria'})
        for t in orden:
            if t not in vistos:
                obs.append((fecha, t, 'turno faltante', ''))
    data.sort(key=lambda r: (r['fecha'], orden[r['turno']]))

    # Desde la primera fecha de la API manda la API; este sitio completa huecos.
    distintos = 0
    if api_path:
        api = {}
        for r in load_workbook(api_path, read_only=True).worksheets[0].iter_rows(min_row=2, values_only=True):
            f = r[0].date() if isinstance(r[0], datetime) else r[0]
            api[(f, r[1])] = {'fecha': f, 'turno': r[1], 'hora': r[3], 'sorteo': '',
                              'numero': int(r[4]), 'mas_1': r[7] or '', 'signo': r[5] or '',
                              'fuente': 'api'}
        desde = min(k[0] for k in api)
        propio = {(r['fecha'], r['turno']): r for r in data}
        for k, v in api.items():
            if k in propio and propio[k]['numero'] != v['numero']:
                distintos += 1
                obs.append((k[0], k[1], 'distinto a la API (se usa la API)',
                            'acá %02d · API %02d' % (propio[k]['numero'], v['numero'])))
        for k, v in propio.items():
            if k[0] >= desde and k not in api:
                obs.append((k[0], k[1], 'hueco de la API completado', '%02d' % v['numero']))
        propio.update(api)
        data = sorted(propio.values(), key=lambda r: (r['fecha'], orden[r['turno']]))
        # Lo que falta en este sitio pero trajo la API ya no es un faltante.
        obs = [o for o in obs if not (o[2] in ('turno faltante', 'fila ilegible',
                                               'fila dudosa (descartada)')
                                      and (o[0], o[1]) in api)]
        dias_api = {k[0] for k in api}
        obs = [o for o in obs if not (o[1] == '' and o[0] in dias_api)]

    wb = Workbook()
    ws = wb.active
    ws.title = 'La Diaria'
    ws.append(['Fecha', 'Turno', 'Turno (nombre)', 'Hora', 'Sorteo nro.', 'Número', 'Más 1',
               'Signo', 'Fuente'])
    for r in data:
        ws.append([r['fecha'], r['turno'], nombre[r['turno']], r['hora'], r['sorteo'],
                   '%02d' % r['numero'], r['mas_1'], r['signo'], r['fuente']])
    for c in ws[1]:
        c.font = Font(bold=True)
    for row in ws.iter_rows(min_row=2, min_col=1, max_col=1):
        row[0].number_format = 'yyyy-mm-dd'
    for row in ws.iter_rows(min_row=2, min_col=6, max_col=6):
        row[0].number_format = '@'
    for i, w in enumerate([12, 11, 15, 10, 11, 9, 7, 16, 15], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = 'A2'

    wo = wb.create_sheet('Observaciones')
    wo.append(['Fecha', 'Turno', 'Observación', 'Detalle'])
    for c in wo[1]:
        c.font = Font(bold=True)
    for o in sorted(obs, key=lambda x: (x[0], orden.get(x[1], 9))):
        wo.append(list(o))
    for row in wo.iter_rows(min_row=2, min_col=1, max_col=1):
        row[0].number_format = 'yyyy-mm-dd'
    for i, w in enumerate([12, 11, 28, 30], 1):
        wo.column_dimensions[get_column_letter(i)].width = w
    wb.save(out_path)

    from collections import Counter
    print('filas', len(data), 'desde', data[0]['fecha'], 'hasta', data[-1]['fecha'])
    print('observaciones', Counter(o[2] for o in obs), 'distintos a la API', distintos)
    print('horarios', Counter((r['turno'], r['hora'].lower()) for r in data))


if __name__ == '__main__':
    main()
