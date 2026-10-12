"""Completa el Excel de La Diaria con lotodehonduras.com (una página por año).

Uso: python completar_con_lotodehonduras.py <entrada.xlsx> <salida.xlsx>

lotodehonduras.com/la-diaria/<año> lista el año entero, día por día
("Resultado 16 Marzo · 11:00 a.m. 48 + 00 Estrella · 3:00 p.m. …"). Tiene
desde 2016. Se usa para:
  · tapar los turnos que faltan en el Excel;
  · resolver las diferencias con nacionalloteria: se descarta el valor que
    repite otro turno del mismo día en su propia fuente (error de copia
    típico); si ninguno o los dos lo hacen, gana lotodehonduras.
Lo que vino de la API (2025 en adelante) no se toca. Solo hasta ayer.
"""
import html
import re
import sys
import time
from datetime import date, datetime, timedelta

import requests
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

URL = 'https://www.lotodehonduras.com/la-diaria/%d'
FIRST_YEAR = 2016
MESES = 'Enero Febrero Marzo Abril Mayo Junio Julio Agosto Septiembre Octubre Noviembre Diciembre'.split()
MES = {m: i + 1 for i, m in enumerate(MESES)}
BLOQUE = re.compile(r'Resultado (\d{1,2}) (' + '|'.join(MESES) + r')\b')
SORTEO = re.compile(r'(\d{1,2}):(\d{2}) ([ap])\.?\s?m\.? (\d{2})(?: \+ (\d{1,2}))? (.*?)(?= \d{1,2}:\d{2} |$)')
ORDEN = {'morning': 0, 'afternoon': 1, 'evening': 2}
# Por dentro se trabaja con morning/afternoon/evening (los usan los otros
# scripts); en el Excel final van los turnos por horario de La Diaria. Los
# sorteos de antes del 11/02/2020 (10:00 y 14:00) van igual al de 11:00 y 15:00:
# la hora real queda en la columna Hora.
CODIGO = {'morning': 'h1100', 'afternoon': 'h1500', 'evening': 'h2100'}
NOMBRE = {'morning': '11:00 AM', 'afternoon': '3:00 PM', 'evening': '9:00 PM'}
INTERNO = {v: k for k, v in CODIGO.items()}


def turno_de(h, ap):
    if ap == 'a':
        return 'morning'
    if h in (12, 1, 2, 3, 4):
        return 'afternoon'
    if h in (8, 9, 10):
        return 'evening'
    return None


def bajar_anio(session, anio):
    r = session.get(URL % anio, timeout=60)
    if r.status_code == 404:
        return {}
    r.raise_for_status()
    s = re.sub(r'(?s)<(script|style).*?</\1>', '', r.text)
    s = re.sub(r'<[^>]+>', ' ', s)
    s = re.sub(r'\s+', ' ', html.unescape(s))
    partes = BLOQUE.split(s)
    salidas = {}
    for i in range(1, len(partes), 3):
        try:
            fecha = date(anio, MES[partes[i + 1]], int(partes[i]))
        except ValueError:
            continue
        cuerpo = partes[i + 2]
        corte = cuerpo.find('Resultados de')       # encabezado del mes siguiente
        if corte >= 0:
            cuerpo = cuerpo[:corte]
        for h, mm, ap, num, mas, signo in SORTEO.findall(cuerpo):
            turno = turno_de(int(h), ap)
            if turno and (fecha, turno) not in salidas:
                salidas[(fecha, turno)] = {
                    'fecha': fecha, 'turno': turno, 'hora': f'{h}:{mm} {ap}m',
                    'sorteo': '', 'numero': int(num),
                    'mas_1': str(int(mas)) if mas and int(mas) <= 9 else '',
                    'signo': signo.strip(), 'fuente': 'lotodehonduras'}
    return salidas


def repite(dic, k):
    """¿El número de k aparece en otro turno del mismo día de esa fuente?"""
    return any(t != k[1] and (k[0], t) in dic and dic[(k[0], t)]['numero'] == dic[k]['numero']
               for t in ORDEN)


def main():
    entrada, salida = sys.argv[1], sys.argv[2]
    wb_in = load_workbook(entrada, read_only=True)
    cols = None
    actual, obs = {}, []
    for i, r in enumerate(wb_in['La Diaria'].iter_rows(values_only=True)):
        if i == 0:
            cols = list(r)
            continue
        f = r[0].date() if isinstance(r[0], datetime) else r[0]
        t = INTERNO.get(r[1], r[1])
        actual[(f, t)] = {'fecha': f, 'turno': t, 'hora': r[3], 'sorteo': r[4] or '',
                             'numero': int(r[5]), 'mas_1': r[6] or '', 'signo': r[7] or '',
                             'fuente': r[8]}
    for r in list(wb_in['Observaciones'].iter_rows(min_row=2, values_only=True)):
        f = r[0].date() if isinstance(r[0], datetime) else r[0]
        obs.append((f, INTERNO.get(r[1] or '', r[1] or ''), r[2], r[3] or ''))

    s = requests.Session()
    s.headers['User-Agent'] = ('Mozilla/5.0 (Windows NT 10.0; Win64; x64) '
                               'AppleWebKit/537.36 (KHTML, like Gecko) Chrome/147.0 Safari/537.36')
    web = {}
    for anio in range(FIRST_YEAR, date.today().year + 1):
        web.update(bajar_anio(s, anio))
        time.sleep(1)
    ayer = date.today() - timedelta(days=1)
    web = {k: v for k, v in web.items() if k[0] <= ayer}

    tapados, cambiados = [], []
    for k, w in sorted(web.items()):
        a = actual.get(k)
        if a is None:
            actual[k] = w
            tapados.append(k)
            obs.append((k[0], k[1], 'hueco completado con lotodehonduras', '%02d' % w['numero']))
        elif a['fuente'] != 'api' and a['numero'] != w['numero']:
            gana_web = not (repite(web, k) and not repite(actual, k))
            elegido = w if gana_web else a
            obs.append((k[0], k[1], 'diferencia entre fuentes',
                        'nacionalloteria %02d · lotodehonduras %02d → %02d'
                        % (a['numero'], w['numero'], elegido['numero'])))
            if gana_web:
                actual[k] = w
                cambiados.append(k)

    # Lo que ahora existe ya no es faltante.
    obs = [o for o in obs if not (o[2] in ('turno faltante', 'fila ilegible',
                                           'fila dudosa (descartada)') and (o[0], o[1]) in actual)]
    dias = {k[0] for k in actual}
    obs = [o for o in obs if not (o[2] in ('sin_datos', 'sin_tabla') and o[0] in dias
                                  and all((o[0], t) in actual for t in ORDEN))]
    for o in list(obs):
        if o[2] in ('sin_datos', 'sin_tabla') and o[0] in dias:
            obs.remove(o)
            obs += [(o[0], t, 'turno faltante', '') for t in ORDEN if (o[0], t) not in actual]

    data = sorted(actual.values(), key=lambda r: (r['fecha'], ORDEN[r['turno']]))
    wb = Workbook()
    ws = wb.active
    ws.title = 'La Diaria'
    ws.append(cols)
    for r in data:
        ws.append([r['fecha'], CODIGO[r['turno']], NOMBRE[r['turno']], r['hora'], r['sorteo'],
                   '%02d' % r['numero'], r['mas_1'], r['signo'], r['fuente']])
    for c in ws[1]:
        c.font = Font(bold=True)
    for row in ws.iter_rows(min_row=2, min_col=1, max_col=1):
        row[0].number_format = 'yyyy-mm-dd'
    for row in ws.iter_rows(min_row=2, min_col=6, max_col=6):
        row[0].number_format = '@'
    for i, w in enumerate([12, 11, 15, 10, 11, 9, 7, 16, 15], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = 'A2'

    wo = wb.create_sheet('Observaciones')
    wo.append(['Fecha', 'Turno', 'Observación', 'Detalle'])
    for c in wo[1]:
        c.font = Font(bold=True)
    for o in sorted(set(obs), key=lambda x: (x[0], ORDEN.get(x[1], 9), x[2])):
        wo.append([o[0], CODIGO.get(o[1], o[1]), o[2], o[3]])
    for row in wo.iter_rows(min_row=2, min_col=1, max_col=1):
        row[0].number_format = 'yyyy-mm-dd'
    for i, w in enumerate([12, 11, 34, 50], 1):
        wo.column_dimensions[get_column_letter(i)].width = w
    wb.save(salida)

    from collections import Counter
    print('filas', len(data), 'tapados', len(tapados), 'cambiados', len(cambiados))
    print('por fuente', Counter(r['fuente'] for r in data))
    print('observaciones', Counter(o[2] for o in set(obs)))


if __name__ == '__main__':
    main()
