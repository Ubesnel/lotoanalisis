"""Baja el histórico completo de La Diaria (Honduras) y lo vuelca a Excel.

Fuente: api.loteriasdehonduras.com (misma plataforma que conectate). Cada
petición devuelve los 8 sorteos que terminan en la fecha pedida; cada turno
es un juego distinto. Los resultados vienen como ids de opciones del
score_layout, que se traducen con ese mismo layout.
"""
import json
import sys
import time
from datetime import date, datetime, timedelta

import requests
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

API = 'https://api.loteriasdehonduras.com/honduras/site-games/%s'
GAMES = [  # (gameId, código de turno en Odoo, nombre, hora HN)
    ('693ae5bbd7b13e9daed23b31', 'morning',   'Mañana', '11:00'),
    ('693ae5bbd7b13e9daed23b07', 'afternoon', 'Tarde',  '15:00'),
    ('693ae5bbd7b13e9daed23b1f', 'evening',   'Noche',  '21:00'),
]
DELAY = 0.4

out_path = sys.argv[1]
s = requests.Session()
s.headers.update({
    'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) Chrome/147.0 Safari/537.36',
    'Accept': 'application/json',
    'Origin': 'https://loteriasdehonduras.com',
    'Referer': 'https://loteriasdehonduras.com/',
})


def fetch(game_id, hasta):
    r = s.get(API % game_id, params={'date': '%sT06:00:00.000Z' % hasta.isoformat()}, timeout=60)
    r.raise_for_status()
    return json.loads(r.text)['game']


rows, issues = {}, []
for gid, turno, nombre, hora in GAMES:
    cursor, calls = date.today(), 0
    while True:
        g = fetch(gid, cursor)
        calls += 1
        lay = g['score_layout'][0]
        labels = [sl['label'] for sl in lay]
        opts = [{o['id']: o['text'] for o in (sl.get('options') or [])} for sl in lay]
        if calls == 1:
            print(turno, g['title'], 'layout', labels, 'schedule', sorted({x['time'] for x in g.get('schedule') or []}))
        ses = g.get('sessions') or []
        if not ses:
            break
        for x in ses:
            d = datetime.strptime(x['date'][:10], '%Y-%m-%d').date()
            sc = (x.get('score') or [[]])[0]
            vals = [opts[i].get(v, v) if i < len(opts) else v for i, v in enumerate(sc)]
            signo = vals[0] if vals else ''
            try:
                num = int(str(signo).split()[0])
                animal = ' '.join(str(signo).split()[1:])
            except (ValueError, IndexError):
                issues.append((d, turno, 'número ilegible', str(sc)))
                continue
            rec = {'fecha': d, 'turno': turno, 'turno_nombre': nombre, 'hora': hora,
                   'numero': num, 'signo': animal,
                   'multi_x': vals[1] if len(vals) > 1 else '',
                   'mas_1': vals[2] if len(vals) > 2 else '',
                   'created': x.get('createdAt', '')}
            k = (d, turno)
            if k in rows and rows[k]['numero'] != num:
                issues.append((d, turno, 'duplicado con distinto número',
                               f"{rows[k]['numero']:02d} vs {num:02d}"))
            rows[k] = rec
        oldest = min(datetime.strptime(x['date'][:10], '%Y-%m-%d').date() for x in ses)
        nuevo = oldest - timedelta(days=1)
        if nuevo >= cursor:
            break
        cursor = nuevo
        time.sleep(DELAY)
    print(turno, 'peticiones', calls, 'sorteos', sum(1 for k in rows if k[1] == turno))

# Huecos: días entre el primero y el último de cada turno sin sorteo.
for _, turno, _, _ in GAMES:
    ds = sorted(k[0] for k in rows if k[1] == turno)
    have = set(ds)
    d = ds[0]
    while d <= ds[-1]:
        if d not in have:
            issues.append((d, turno, 'sin sorteo en la fuente', ''))
        d += timedelta(days=1)

orden = {'morning': 0, 'afternoon': 1, 'evening': 2}
data = sorted(rows.values(), key=lambda r: (r['fecha'], orden[r['turno']]))

wb = Workbook()
ws = wb.active
ws.title = 'La Diaria'
head = ['Fecha', 'Turno', 'Turno (nombre)', 'Hora HN', 'Número', 'Signo', 'Multi X', 'Más 1']
ws.append(head)
for r in data:
    ws.append([r['fecha'], r['turno'], r['turno_nombre'], r['hora'], '%02d' % r['numero'],
               r['signo'], r['multi_x'], r['mas_1']])
for c in ws[1]:
    c.font = Font(bold=True)
for row in ws.iter_rows(min_row=2, min_col=1, max_col=1):
    row[0].number_format = 'yyyy-mm-dd'
for row in ws.iter_rows(min_row=2, min_col=5, max_col=5):
    row[0].number_format = '@'
for i, w in enumerate([12, 11, 15, 9, 9, 16, 9, 7], 1):
    ws.column_dimensions[get_column_letter(i)].width = w
ws.freeze_panes = 'A2'

wi = wb.create_sheet('Observaciones')
wi.append(['Fecha', 'Turno', 'Observación', 'Detalle'])
for c in wi[1]:
    c.font = Font(bold=True)
for d, t, o, det in sorted(issues, key=lambda x: (x[0], orden.get(x[1], 9))):
    wi.append([d, t, o, det])
for row in wi.iter_rows(min_row=2, min_col=1, max_col=1):
    row[0].number_format = 'yyyy-mm-dd'
for i, w in enumerate([12, 11, 30, 20], 1):
    wi.column_dimensions[get_column_letter(i)].width = w

wb.save(out_path)
print('filas', len(data), 'desde', data[0]['fecha'], 'hasta', data[-1]['fecha'], 'observaciones', len(issues))
from collections import Counter
print(Counter(o for _, _, o, _ in issues))
# Muestra para validar fecha vs hora real de carga (createdAt en UTC).
for r in data[-6:]:
    print(r['fecha'], r['turno'], '%02d' % r['numero'], r['signo'], r['created'])
