"""Histórico de Maryland, South Carolina e Illinois (Pick 3 + Pick 4) a Excel, listo para el wizard.

Uso: python extract_md_sc_il.py <maryland|south_carolina|illinois> <cache.json> <salida.xlsx>

Fuente: lottery.net, una página por juego, turno y año, igual que New Jersey:
  https://www.lottery.net/<estado>/pick-3-midday/numbers/<año>   (→ Tarde)
  https://www.lottery.net/<estado>/pick-3-evening/numbers/<año>  (→ Noche)
y lo mismo con pick-4. Pick 3 da la centena y el número (últimos dos dígitos);
Pick 4 da los premios 2 y 3 (primer y segundo par). La Fireball se ignora.
Maryland e Illinois desde abril de 2010;
South Carolina desde 2003 (la Tarde desde el 01/07/2003 y nunca los domingos).
El año en curso se vuelve a pedir siempre; los cerrados quedan en el caché JSON.
"""
import json
import os
import re
import sys
import time
from collections import Counter
from datetime import date, timedelta

import requests
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

URL = 'https://www.lottery.net/%s/pick-%d-%s/numbers/%d'
ESTADOS = {
    # clave: (slug en lottery.net, título de la hoja, primer año, días sin Tarde (0=lunes))
    'maryland': ('maryland', 'Maryland', 2010, ()),
    'south_carolina': ('south-carolina', 'South Carolina', 2003, (6,)),
    'illinois': ('illinois', 'Chicago - Illinois', 2010, ()),
}
TURNOS = {'midday': ('afternoon', 'Tarde'), 'evening': ('evening', 'Noche')}
DELAY = 2.0

RE_FILA = re.compile(r'<tr>\s*<td class="colour noBefore pickArchive">\s*([^<]+?)\s*</td>(.*?)</tr>', re.S)
# Maryland Evening sale con la clase nj-pick-3: no se filtra por estado.
RE_BOLA = re.compile(r'class="ball [a-z]+-pick-[34] number-part-0([1-4])">(\d)<')
MESES = {m: i for i, m in enumerate(
    ['January', 'February', 'March', 'April', 'May', 'June', 'July', 'August',
     'September', 'October', 'November', 'December'], 1)}


def parsear(html, bolas_n):
    """→ {iso_fecha: 'ABC'} con las bolas principales (3 o 4 dígitos)."""
    out = {}
    for fecha_txt, cuerpo in RE_FILA.findall(html):
        partes = fecha_txt.replace(',', '').split()
        try:
            d = date(int(partes[3]), MESES[partes[1]], int(partes[2]))
        except (ValueError, IndexError, KeyError):
            continue
        bolas = dict(RE_BOLA.findall(cuerpo))
        if len(bolas) == bolas_n:
            out[d.isoformat()] = ''.join(bolas[str(i)] for i in range(1, bolas_n + 1))
    return out


def descargar(slug, primer_anio, cache_path, anio_actual):
    cache = {}
    if os.path.exists(cache_path):
        with open(cache_path, encoding='utf-8') as f:
            cache = json.load(f)
    s = requests.Session()
    s.headers['User-Agent'] = ('Mozilla/5.0 (Windows NT 10.0; Win64; x64) '
                               'AppleWebKit/537.36 (KHTML, like Gecko) Chrome/147.0 Safari/537.36')
    for anio in range(primer_anio, anio_actual + 1):
        for juego in (3, 4):
            for turno in TURNOS:
                k = 'p%d/%s/%d' % (juego, turno, anio)
                if k in cache and anio < anio_actual:
                    continue
                for intento in range(3):
                    try:
                        r = s.get(URL % (slug, juego, turno, anio), timeout=60)
                        r.raise_for_status()
                        cache[k] = parsear(r.text, juego)
                        print(k, len(cache[k]))
                        break
                    except Exception as exc:
                        print('error', k, exc)
                        time.sleep(5 * (intento + 1))
                time.sleep(DELAY)
        # Se guarda por año para no perder lo bajado si se corta.
        with open(cache_path, 'w', encoding='utf-8') as f:
            json.dump(cache, f)
    return cache


def main():
    estado, cache_path, salida = sys.argv[1], sys.argv[2], sys.argv[3]
    slug, titulo, primer_anio, sin_tarde = ESTADOS[estado]
    hoy = date.today()
    cache = descargar(slug, primer_anio, cache_path, hoy.year)

    filas, obs = [], []
    for turno, (codigo, nombre) in TURNOS.items():
        datos, datos4 = {}, {}
        for k, v in cache.items():
            if k.startswith('p3/%s/' % turno):
                datos.update(v)
            elif k.startswith('p4/%s/' % turno):
                datos4.update(v)
        if not datos:
            continue
        # Arranca en el primer sorteo que tiene la fuente para ese turno.
        d = date.fromisoformat(min(datos))
        while d <= hoy:
            k = d.isoformat()
            if turno == 'midday' and d.weekday() in sin_tarde:
                pass
            elif k in datos and k in datos4:
                filas.append((d, codigo, nombre, datos[k], datos4[k]))
            elif d < hoy:
                faltan = 'Pick 3' if k not in datos else 'Pick 4'
                if k not in datos and k not in datos4:
                    faltan = 'Pick 3 y Pick 4'
                obs.append((d, codigo, 'sin resultado en la fuente (se omite el turno)', faltan))
            d += timedelta(days=1)

    filas.sort(key=lambda r: (r[0], r[1] != 'afternoon'))
    wb = Workbook()
    ws = wb.active
    ws.title = titulo
    ws.append(['Fecha', 'Turno', 'Turno (nombre)', 'Pick 3', 'Pick 4', 'Centena', 'Número',
               'Premio 2', 'Premio 3', 'Fuente'])
    for fecha, codigo, nombre, p3, p4 in filas:
        ws.append([fecha, codigo, nombre, p3, p4, int(p3[0]), p3[1:], p4[:2], p4[2:], 'lottery.net'])
    for c in ws[1]:
        c.font = Font(bold=True)
    for row in ws.iter_rows(min_row=2, min_col=1, max_col=1):
        row[0].number_format = 'yyyy-mm-dd'
    for col in (4, 5, 7, 8, 9):
        for row in ws.iter_rows(min_row=2, min_col=col, max_col=col):
            row[0].number_format = '@'
    for i, w in enumerate([12, 10, 15, 8, 8, 9, 9, 9, 9, 12], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = 'A2'

    wo = wb.create_sheet('Observaciones')
    wo.append(['Fecha', 'Turno', 'Observación', 'Detalle'])
    for c in wo[1]:
        c.font = Font(bold=True)
    for o in sorted(obs, key=lambda x: (x[0], x[1])):
        wo.append(list(o))
    for row in wo.iter_rows(min_row=2, min_col=1, max_col=1):
        row[0].number_format = 'yyyy-mm-dd'
    for i, w in enumerate([12, 10, 46, 16], 1):
        wo.column_dimensions[get_column_letter(i)].width = w
    wb.save(salida)

    print('filas', len(filas), 'desde', filas[0][0], 'hasta', filas[-1][0])
    print('por turno', Counter(r[1] for r in filas))
    print('observaciones', len(obs), Counter(o[0].year for o in obs))


if __name__ == '__main__':
    main()
