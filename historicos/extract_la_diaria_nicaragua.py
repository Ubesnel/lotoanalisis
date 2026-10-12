"""Histórico de La Diaria de Nicaragua a Excel, listo para el wizard.

Uso: python extract_la_diaria_nicaragua.py <salida.xlsx>

Dos fuentes:
  · La Nueva Radio YA (nuevaya.com.ni): una página por año con todos los
    días, desde el 01/01/2024. Es la base, porque tiene todos los turnos.
  · API de loteriasdenicaragua.com (misma plataforma que la de Honduras):
    desde el 01/06/2025. Cada resultado lo cargan dos operadores y tiene que
    coincidir, así que en una diferencia gana la API. Le faltan casi todos los
    sorteos de 6:00 PM que no son de sábado.

Horarios (hora de Nicaragua), según los datos:
  · hasta el 24/08/2025: 11, 15, 21 todos los días; 18 los sábados.
  · 25/08/2025 → 14/06/2026: 4 por día; el primero pasó a las 12:00 desde
    el 01/09/2025.
  · desde el 15/06/2026: 11, 15, 21 todos los días; 18 sábado y domingo.
El sorteo de 12:00 va al turno de 11:00 AM (es el primero del día); la hora
real queda en la columna Hora.
"""
import html
import json
import re
import sys
import time
from collections import Counter
from datetime import date, datetime, timedelta

import requests
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

YA_PAGES = {
    2024: 'https://nuevaya.com.ni/loto-diaria-de-nicaragua-resultados-anteriores-ano-2024/',
    2025: 'https://nuevaya.com.ni/loto-diaria-de-nicaragua-resultados-anteriores-ano-2025/',
    2026: 'https://nuevaya.com.ni/historico-loto-nicaragua-2026-resultados-por-mes/',
}
API = 'https://api.loteriasdenicaragua.com/nicaragua/site-games/%s'
API_GAMES = {  # gameId → turno
    '6938bae65aada821ed601e99': 'h1100',   # "Diaria 11:00" (trae también los de 12:00)
    '6938bae65aada821ed601ebd': 'h1500',
    '6938bae65aada821ed601f00': 'h1800',   # solo trae sábados
    '6938bae65aada821ed601ed6': 'h2100',
}
TURNO_DE_HORA = {11: 'h1100', 12: 'h1100', 15: 'h1500', 18: 'h1800', 21: 'h2100'}
NOMBRE = {'h1100': '11:00 AM', 'h1500': '3:00 PM', 'h1800': '6:00 PM', 'h2100': '9:00 PM'}
ORDEN = {t: i for i, t in enumerate(NOMBRE)}
MESES = 'Enero Febrero Marzo Abril Mayo Junio Julio Agosto Septiembre Octubre Noviembre Diciembre'.split()
MES = {m: i + 1 for i, m in enumerate(MESES)}
DELAY = 0.5


def sesion(**headers):
    s = requests.Session()
    s.headers['User-Agent'] = ('Mozilla/5.0 (Windows NT 10.0; Win64; x64) '
                               'AppleWebKit/537.36 (KHTML, like Gecko) Chrome/147.0 Safari/537.36')
    s.headers.update(headers)
    return s


def hora_de(texto):
    """'11:00 am' / '12:00 m' / '6.00 pm' → hora 0-23, o None."""
    m = re.match(r'(\d{1,2})[:.](\d{2})\s*(am|pm|m)\b', texto.strip().lower())
    if not m:
        return None
    h = int(m.group(1))
    if m.group(3) == 'pm' and h != 12:
        h += 12
    return h


def bajar_ya(obs):
    s = sesion()
    salidas = {}
    for anio, url in YA_PAGES.items():
        r = s.get(url, timeout=60)
        r.raise_for_status()
        t = re.sub(r'(?s)<(script|style).*?</\1>', '', r.text)
        t = re.sub(r'<[^>]+>', ' | ', t)
        t = re.sub(r'(\s*\|\s*)+', ' | ', html.unescape(t))
        t = re.sub(r'\s+', ' ', t)
        for dd, mm, cuerpo in re.findall(
                r'(\d{1,2}) de (' + '|'.join(MESES) + r') \| Hora \| Resultados \| Diaria \| (.*?) \| Fechas', t):
            try:
                fecha = date(anio, MES[mm], int(dd))
            except ValueError:
                obs.append((None, '', 'fecha inválida en YA', f'{dd} de {mm} de {anio}'))
                continue
            # En los días recientes la tabla de Jugá 4 viene pegada a la de Diaria.
            cuerpo = cuerpo.split(' | Jugá 4')[0]
            celdas = cuerpo.split(' | ')
            for i in range(0, len(celdas) - 1, 2):
                h = hora_de(celdas[i])
                m = re.match(r'(\d{2})\b(.*)$', celdas[i + 1].strip())
                if h not in TURNO_DE_HORA or not m:
                    obs.append((fecha, '', 'fila ilegible en YA', f'{celdas[i]} | {celdas[i + 1]}'))
                    continue
                turno = TURNO_DE_HORA[h]
                resto = m.group(2)
                multi = re.search(r'\((JG|\d+X)\)', resto, re.I)
                mas = re.search(r'\+1\s*=\s*(\d)', resto)
                salidas.setdefault((fecha, turno), {
                    'fecha': fecha, 'turno': turno, 'hora': '%02d:00' % h,
                    'numero': int(m.group(1)), 'signo': '',
                    'multi_x': multi.group(1).upper() if multi else '',
                    'mas_1': mas.group(1) if mas else '', 'fuente': 'nuevaya'})
        time.sleep(1)
    return salidas


def bajar_api():
    s = sesion(Origin='https://loteriasdenicaragua.com', Referer='https://loteriasdenicaragua.com/')
    salidas = {}
    for gid, turno in API_GAMES.items():
        cursor = date.today()
        while True:
            r = s.get(API % gid, params={'date': cursor.isoformat() + 'T06:00:00.000Z'}, timeout=60)
            r.raise_for_status()
            g = json.loads(r.text)['game']
            lay = g['score_layout'][0]
            etiquetas = [sl['label'] for sl in lay]
            opciones = {o['id']: o['text'] for sl in lay for o in (sl.get('options') or [])}
            sesiones = g.get('sessions') or []
            if not sesiones:
                break
            for x in sesiones:
                fecha = datetime.strptime(x['date'][:10], '%Y-%m-%d').date()
                vals = dict(zip(etiquetas, (opciones.get(v, v) for v in x['score'][0])))
                signo = str(vals.get('Signo', ''))
                if not re.match(r'\d{2}\b', signo):
                    continue
                salidas[(fecha, turno)] = {
                    'fecha': fecha, 'turno': turno, 'hora': '',
                    'numero': int(signo[:2]), 'signo': signo[2:].strip(),
                    'multi_x': vals.get('Multi X', ''), 'mas_1': vals.get('Más 1', ''),
                    'fuente': 'api'}
            viejo = min(datetime.strptime(x['date'][:10], '%Y-%m-%d').date() for x in sesiones)
            if viejo - timedelta(days=1) >= cursor:
                break
            cursor = viejo - timedelta(days=1)
            time.sleep(DELAY)
    return salidas


def esperados(fecha):
    """Turnos que se sortearon ese día según el esquema vigente."""
    if fecha < date(2025, 8, 25):
        return ['h1100', 'h1500', 'h2100'] + (['h1800'] if fecha.weekday() == 5 else [])
    if fecha <= date(2026, 6, 14):
        return ['h1100', 'h1500', 'h1800', 'h2100']
    return ['h1100', 'h1500', 'h2100'] + (['h1800'] if fecha.weekday() >= 5 else [])


def main():
    salida = sys.argv[1]
    obs = []
    ya = bajar_ya(obs)
    api = bajar_api()
    hasta = date.today() - timedelta(days=1)

    final = {}
    for k, v in ya.items():
        if k[0] <= hasta:
            final[k] = v
    for k, a in api.items():
        if k[0] > hasta:
            continue
        y = ya.get(k)
        if y and y['numero'] != a['numero']:
            obs.append((k[0], k[1], 'diferencia entre fuentes (se usa la API)',
                        'nuevaya %02d · API %02d' % (y['numero'], a['numero'])))
        if not y:
            obs.append((k[0], k[1], 'falta en nuevaya, completado con la API', '%02d' % a['numero']))
        # La hora real la sabe nuevaya (la API mete el de 12:00 en su juego de 11:00).
        a = dict(a, hora=y['hora'] if y else NOMBRE[k[1]][:5].strip())
        final[k] = a

    d = min(k[0] for k in final)
    while d <= hasta:
        for t in esperados(d):
            if (d, t) not in final:
                obs.append((d, t, 'turno faltante', ''))
        d += timedelta(days=1)

    data = sorted(final.values(), key=lambda r: (r['fecha'], ORDEN[r['turno']]))
    wb = Workbook()
    ws = wb.active
    ws.title = 'La Diaria'
    ws.append(['Fecha', 'Turno', 'Turno (nombre)', 'Hora', 'Número', 'Signo', 'Multi X', 'Más 1', 'Fuente'])
    for r in data:
        ws.append([r['fecha'], r['turno'], NOMBRE[r['turno']], r['hora'], '%02d' % r['numero'],
                   r['signo'], r['multi_x'], r['mas_1'], r['fuente']])
    for c in ws[1]:
        c.font = Font(bold=True)
    for row in ws.iter_rows(min_row=2, min_col=1, max_col=1):
        row[0].number_format = 'yyyy-mm-dd'
    for row in ws.iter_rows(min_row=2, min_col=5, max_col=5):
        row[0].number_format = '@'
    for i, w in enumerate([12, 9, 15, 8, 9, 16, 9, 7, 10], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = 'A2'

    wo = wb.create_sheet('Observaciones')
    wo.append(['Fecha', 'Turno', 'Observación', 'Detalle'])
    for c in wo[1]:
        c.font = Font(bold=True)
    for o in sorted(obs, key=lambda x: (x[0] or date.min, ORDEN.get(x[1], 9))):
        wo.append(list(o))
    for row in wo.iter_rows(min_row=2, min_col=1, max_col=1):
        row[0].number_format = 'yyyy-mm-dd'
    for i, w in enumerate([12, 9, 40, 30], 1):
        wo.column_dimensions[get_column_letter(i)].width = w
    wb.save(salida)

    print('filas', len(data), 'desde', data[0]['fecha'], 'hasta', data[-1]['fecha'])
    print('por fuente', Counter(r['fuente'] for r in data))
    print('por turno', Counter(r['turno'] for r in data))
    print('observaciones', Counter(o[2] for o in obs))


if __name__ == '__main__':
    main()
