"""Histórico de New Jersey (Pick-3 + Pick-4) a Excel, listo para el wizard.

Uso: python extract_new_jersey.py <cache.json> <salida.xlsx>

Fuente: lottery.net, una página por turno y año:
  https://www.lottery.net/new-jersey/pick-3-midday/numbers/<año>   (12:59 ET → Tarde)
  https://www.lottery.net/new-jersey/pick-3-evening/numbers/<año>  (22:57 ET → Noche)
y lo mismo con pick-4-midday / pick-4-evening. Pick-3 da la centena y el número
(últimos dos dígitos); Pick-4 da los premios 2 y 3 (primer y segundo par de
dígitos, como en New York y Florida). La Fireball se ignora. Desde 2010. El año en curso se vuelve a pedir
siempre; los cerrados quedan en el caché JSON.
"""
import json
import os
import re
import sys
import time
from collections import Counter
from datetime import date, datetime, timedelta

import requests
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

URL = 'https://www.lottery.net/new-jersey/pick-%d-%s/numbers/%d'
PRIMER_ANIO = 2010
TURNOS = {'midday': ('afternoon', 'Tarde'), 'evening': ('evening', 'Noche')}
DELAY = 1.0

RE_FILA = re.compile(r'<tr>\s*<td class="colour noBefore pickArchive">\s*([^<]+?)\s*</td>(.*?)</tr>', re.S)
RE_BOLA = re.compile(r'class="ball nj-pick-[34] number-part-0([1-4])">(\d)<')


def parsear(html, bolas_n):
    """→ {iso_fecha: 'ABC'} con las bolas principales (3 o 4 dígitos)."""
    out = {}
    for fecha_txt, cuerpo in RE_FILA.findall(html):
        try:
            d = datetime.strptime(fecha_txt.split(' ', 1)[1].strip(), '%B %d, %Y').date()
        except (ValueError, IndexError):
            continue
        bolas = dict(RE_BOLA.findall(cuerpo))
        if len(bolas) == bolas_n:
            out[d.isoformat()] = ''.join(bolas[str(i)] for i in range(1, bolas_n + 1))
    return out


def descargar(cache_path, anio_actual):
    cache = {}
    if os.path.exists(cache_path):
        with open(cache_path, encoding='utf-8') as f:
            cache = json.load(f)
    s = requests.Session()
    s.headers['User-Agent'] = ('Mozilla/5.0 (Windows NT 10.0; Win64; x64) '
                               'AppleWebKit/537.36 (KHTML, like Gecko) Chrome/147.0 Safari/537.36')
    for anio in range(PRIMER_ANIO, anio_actual + 1):
        for juego in (3, 4):
            for turno in TURNOS:
                k = 'p%d/%s/%d' % (juego, turno, anio)
                if k in cache and anio < anio_actual:
                    continue
                for intento in range(3):
                    try:
                        r = s.get(URL % (juego, turno, anio), timeout=60)
                        r.raise_for_status()
                        cache[k] = parsear(r.text, juego)
                        break
                    except Exception as exc:
                        print('error', k, exc)
                        time.sleep(5 * (intento + 1))
                time.sleep(DELAY)
    with open(cache_path, 'w', encoding='utf-8') as f:
        json.dump(cache, f)
    return cache


def main():
    cache_path, salida = sys.argv[1], sys.argv[2]
    hoy = date.today()
    cache = descargar(cache_path, hoy.year)

    filas, obs = [], []
    for turno, (codigo, nombre) in TURNOS.items():
        datos, datos4 = {}, {}
        for k, v in cache.items():
            if k.startswith('p3/%s/' % turno):
                datos.update(v)
            elif k.startswith('p4/%s/' % turno):
                datos4.update(v)
        d = date(PRIMER_ANIO, 1, 1)
        while d <= hoy:
            k = d.isoformat()
            if k in datos and k in datos4:
                filas.append((d, codigo, nombre, datos[k], datos4[k]))
            elif d < hoy:
                faltan = 'Pick 3' if k not in datos else 'Pick 4'
                if k not in datos and k not in datos4:
                    faltan = 'Pick 3 y Pick 4'
                obs.append((d, codigo, 'sin resultado en la fuente (se omite el turno)', faltan))
            d += timedelta(days=1)

    filas.sort(key=lambda r: (r[0], r[1] != 'afternoon'))
    wb = Workbook()
    ws = wb.active
    ws.title = 'New Jersey'
    ws.append(['Fecha', 'Turno', 'Turno (nombre)', 'Pick 3', 'Pick 4', 'Centena', 'Número',
               'Premio 2', 'Premio 3', 'Fuente'])
    for fecha, codigo, nombre, p3, p4 in filas:
        ws.append([fecha, codigo, nombre, p3, p4, int(p3[0]), p3[1:], p4[:2], p4[2:], 'lottery.net'])
    for c in ws[1]:
        c.font = Font(bold=True)
    for row in ws.iter_rows(min_row=2, min_col=1, max_col=1):
        row[0].number_format = 'yyyy-mm-dd'
    for col in (4, 5, 7, 8, 9):
        for row in ws.iter_rows(min_row=2, min_col=col, max_col=col):
            row[0].number_format = '@'
    for i, w in enumerate([12, 10, 15, 8, 8, 9, 9, 9, 9, 12], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = 'A2'

    wo = wb.create_sheet('Observaciones')
    wo.append(['Fecha', 'Turno', 'Observación', 'Detalle'])
    for c in wo[1]:
        c.font = Font(bold=True)
    for o in sorted(obs, key=lambda x: (x[0], x[1])):
        wo.append(list(o))
    for row in wo.iter_rows(min_row=2, min_col=1, max_col=1):
        row[0].number_format = 'yyyy-mm-dd'
    for i, w in enumerate([12, 10, 46, 16], 1):
        wo.column_dimensions[get_column_letter(i)].width = w
    wb.save(salida)

    print('filas', len(filas), 'desde', filas[0][0], 'hasta', filas[-1][0])
    print('por turno', Counter(r[1] for r in filas))
    print('observaciones', len(obs), Counter(o[0].year for o in obs))


if __name__ == '__main__':
    main()
