"""Histórico de Nuevos Tiempos Reventados (JPS, Costa Rica) a Excel, listo para el wizard.

Uso: python extract_nuevos_tiempos_cr.py <cache.json> <salida.xlsx> [--sin-bajar]

Fuente: nacionalloteria.com (no oficial). jps.go.cr y su API (integration.jps.go.cr)
bloquean por WAF todo acceso desde fuera de Costa Rica, así que no se usan.
  GET https://www.nacionalloteria.com/costarica/nuevos-tiempos.php?del-dia=YYYY-MM-DD
Una página por día, con tabla Medio día / Tarde / Noche: número, "Rev." (bolita
reventada, etiqueta de color) y "MgRev." (número del Mega Reventado). Un día sin
datos muestra "Aún no está disponibles los premios". Los datos arrancan a fines
de 2014; la Tarde aparece después (antes figura como "--").

El caché guarda solo el bloque de resultados de cada página (crudo, en HTML).
Se baja de lo más nuevo hacia atrás, a DELAY segundos por página, y se corta
solo ante un 403 (nacionalloteria ya bloqueó una vez con las quinielas de RD).
"""
import html
import json
import os
import re
import sys
import time
from collections import Counter
from datetime import date, timedelta

import requests
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

URL = 'https://www.nacionalloteria.com/costarica/nuevos-tiempos.php'
DESDE = date(2014, 9, 1)
REPEDIR_DIAS = 3
DELAY = 3
GUARDAR_CADA = 50

TURNO = {'medio dia': 'h1255', 'tarde': 'h1630', 'noche': 'h1930'}
NOMBRE = {'h1255': '12:55 PM', 'h1630': '4:30 PM', 'h1930': '7:30 PM'}
ORDEN = {t: i for i, t in enumerate(NOMBRE)}

RE_FILA = re.compile(
    r'label-normal-sm-mv">([^<]+)</span></td>\s*'
    r'<td[^>]*><span class="label label-numero">([^<]*)</span>\s*</td>\s*'
    r'<td[^>]*>(.*?)</td>\s*'
    r'<td[^>]*>(?:<span class="label label-numero">([^<]*)</span>)?\s*</td>', re.S)
RE_REV = re.compile(r'<span class="label label-([a-z]+)">([^<]*)</span>')


def _norm(texto):
    texto = html.unescape(texto).strip().lower()
    for a, b in (('á', 'a'), ('é', 'e'), ('í', 'i'), ('ó', 'o'), ('ú', 'u')):
        texto = texto.replace(a, b)
    return texto


def _bloque(pagina):
    """Bloque de resultados de la página: lo que se guarda en el caché."""
    i = pagina.find('<section id="listaResultados">')
    if i < 0:
        return None
    if 'no est' in pagina[i:i + 1500] and 'disponibles' in pagina[i:i + 1500]:
        return ''
    j = pagina.find('</table>', i)
    return pagina[i:j + 8] if j > 0 else None


def descargar(cache_path, hasta):
    cache = {}
    if os.path.exists(cache_path):
        with open(cache_path, encoding='utf-8') as f:
            cache = json.load(f)

    def guardar():
        with open(cache_path, 'w', encoding='utf-8') as f:
            json.dump(cache, f, ensure_ascii=False)

    s = requests.Session()
    s.headers.update({'User-Agent': ('Mozilla/5.0 (Windows NT 10.0; Win64; x64) '
                                     'AppleWebKit/537.36 (KHTML, like Gecko) Chrome/147.0 Safari/537.36'),
                      'Referer': URL})
    repedir = hasta - timedelta(days=REPEDIR_DIAS)
    d, pedidas = hasta, 0
    while d >= DESDE:
        k = d.isoformat()
        if k not in cache or 'error' in cache[k] or d > repedir:
            for intento in range(3):
                try:
                    r = s.get(URL, params={'del-dia': k}, timeout=40)
                    if r.status_code == 403:
                        guardar()
                        print('403 en', k, '- se corta; el caché queda guardado')
                        return cache
                    r.raise_for_status()
                    bloque = _bloque(r.content.decode('utf-8', 'replace'))
                    cache[k] = {'error': 'página sin bloque de resultados'} if bloque is None else {'html': bloque}
                    break
                except Exception as exc:
                    cache[k] = {'error': str(exc)}
                    time.sleep(5 * (intento + 1))
            pedidas += 1
            if pedidas % GUARDAR_CADA == 0:
                guardar()
                print('van', pedidas, 'páginas, por', k, flush=True)
            time.sleep(DELAY)
        d -= timedelta(days=1)
    guardar()
    return cache


def parsear(k, dia):
    """[(turno, numero, reventado, mega)] y observaciones de un día del caché."""
    salidas, obs = [], []
    bloque = dia.get('html') or ''
    if not bloque:
        return salidas, obs
    m = re.search(r'<time datetime="(\d{4}-\d{2}-\d{2})"', bloque)
    if m and m.group(1) != k:
        obs.append(('', 'la página muestra otra fecha (descartada)', m.group(1)))
        return salidas, obs
    for nombre, num, rev, mega in RE_FILA.findall(bloque):
        turno = TURNO.get(_norm(nombre))
        if not turno:
            obs.append(('', 'fila desconocida', nombre))
            continue
        num = num.strip()
        if num in ('', '--', '—', '0-'):   # la fuente marca así los sorteos sin dato
            continue
        if not num.isdigit() or int(num) > 99:
            obs.append((turno, 'número ilegible', num))
            continue
        rv = RE_REV.search(rev)
        reventado = '%s:%s' % (rv.group(1), rv.group(2).strip()) if rv else ''
        mega = (mega or '').strip()
        salidas.append((turno, int(num), reventado, mega if mega.isdigit() else ''))
    return salidas, obs


def main():
    cache_path, salida = sys.argv[1], sys.argv[2]
    hasta = date.today() - timedelta(days=1)
    if '--sin-bajar' in sys.argv:
        with open(cache_path, encoding='utf-8') as f:
            cache = json.load(f)
    else:
        cache = descargar(cache_path, hasta)

    data, obs = {}, []
    inicio = {}   # primer día con cada turno
    for k in sorted(cache):
        dia = cache[k]
        d = date.fromisoformat(k)
        if 'error' in dia:
            obs.append((d, '', 'error de descarga', dia['error']))
            continue
        salidas, o = parsear(k, dia)
        obs += [(d,) + x for x in o]
        for turno, num, rev, mega in salidas:
            inicio.setdefault(turno, d)
            data[(d, turno)] = {'fecha': d, 'turno': turno, 'numero': num,
                                'reventado': rev, 'mega': mega}

    # Turnos faltantes: entre el primer día que aparece cada turno y el último dato.
    if data:
        ultimo = max(r['fecha'] for r in data.values())
        for turno, desde in inicio.items():
            d = desde
            while d <= ultimo:
                if (d, turno) not in data:
                    detalle = 'día sin datos' if not any((d, t) in data for t in NOMBRE) else ''
                    obs.append((d, turno, 'turno faltante', detalle))
                d += timedelta(days=1)

    filas = sorted(data.values(), key=lambda r: (r['fecha'], ORDEN[r['turno']]))
    wb = Workbook()
    ws = wb.active
    ws.title = 'Nuevos Tiempos'
    ws.append(['Fecha', 'Turno', 'Turno (nombre)', 'Número', 'Reventado', 'Mega Reventado', 'Fuente'])
    for r in filas:
        ws.append([r['fecha'], r['turno'], NOMBRE[r['turno']], '%02d' % r['numero'],
                   r['reventado'], r['mega'], 'nacionalloteria'])
    for c in ws[1]:
        c.font = Font(bold=True)
    for row in ws.iter_rows(min_row=2, min_col=1, max_col=1):
        row[0].number_format = 'yyyy-mm-dd'
    for row in ws.iter_rows(min_row=2, min_col=4, max_col=6):
        for c in row:
            c.number_format = '@'
    for i, w in enumerate([12, 9, 15, 9, 12, 15, 15], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = 'A2'

    wo = wb.create_sheet('Observaciones')
    wo.append(['Fecha', 'Turno', 'Observación', 'Detalle'])
    for c in wo[1]:
        c.font = Font(bold=True)
    for o in sorted(obs, key=lambda x: (x[0], ORDEN.get(x[1], 9))):
        wo.append(list(o))
    for row in wo.iter_rows(min_row=2, min_col=1, max_col=1):
        row[0].number_format = 'yyyy-mm-dd'
    for i, w in enumerate([12, 9, 42, 16], 1):
        wo.column_dimensions[get_column_letter(i)].width = w
    wb.save(salida)

    print('filas', len(filas), 'desde', filas[0]['fecha'], 'hasta', filas[-1]['fecha'])
    print('por turno', Counter(r['turno'] for r in filas))
    print('inicio de cada turno', {t: str(d) for t, d in inicio.items()})
    print('reventado', Counter(r['reventado'] for r in filas))
    print('observaciones', Counter(o[2] for o in obs))


if __name__ == '__main__':
    main()
