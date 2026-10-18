"""Une gana_mas.xlsx (Tarde) y loteria_nacional.xlsx (Noche), generados por
extract_quinielas_rd.py, en un solo loteria_nacional.xlsx con los dos turnos.
Uso: python unir_loteria_nacional.py <carpeta_con_los_dos_xlsx> <salida.xlsx>"""
import json
import sys
from datetime import datetime
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

carpeta, salida = sys.argv[1], sys.argv[2]
rows, obs = [], []
for f, t in [('gana_mas', 'Tarde'), ('loteria_nacional', 'Noche')]:
    wb = load_workbook(f'{carpeta}/{f}.xlsx')
    rows += list(wb['Salidas'].iter_rows(min_row=2, values_only=True))
    obs += [(o[0], t, o[1], o[2]) for o in wb['Observaciones'].iter_rows(min_row=2, values_only=True)]
# El extractor corta en ayer: si conectate ya trae el Gana Más de hoy, se suma.
cg = json.load(open('cache_quinielas_rd/gana_mas_conectate.json'))
have = {(r[0].date().isoformat(), r[1]) for r in rows}
for k, n in cg.items():
    if (k, 'afternoon') not in have and k > max(r[0] for r in rows).date().isoformat():
        rows.append((datetime.fromisoformat(k), 'afternoon', 'Tarde', *[x.zfill(2) for x in n], 'conectate'))
rows.sort(key=lambda r: (r[0], r[1] != 'afternoon'))
wb = Workbook()
ws = wb.active
ws.title = 'Salidas'
ws.append(['Fecha', 'Turno', 'Turno (nombre)', 'Número', 'Premio 2', 'Premio 3', 'Fuente'])
for r in rows:
    ws.append(list(r))
for c in ws[1]:
    c.font = Font(bold=True)
for row in ws.iter_rows(min_row=2, max_col=1):
    row[0].number_format = 'yyyy-mm-dd'
for row in ws.iter_rows(min_row=2, min_col=4, max_col=6):
    for c in row:
        c.number_format = '@'
for i, w in enumerate([12, 10, 15, 9, 9, 9, 16], 1):
    ws.column_dimensions[get_column_letter(i)].width = w
ws.freeze_panes = 'A2'
wo = wb.create_sheet('Observaciones')
wo.append(['Fecha', 'Turno', 'Observación', 'Detalle'])
for c in wo[1]:
    c.font = Font(bold=True)
for o in sorted(obs, key=lambda o: (o[0], o[1])):
    wo.append(list(o))
for row in wo.iter_rows(min_row=2, max_col=1):
    row[0].number_format = 'yyyy-mm-dd'
for i, w in enumerate([12, 8, 44, 40], 1):
    wo.column_dimensions[get_column_letter(i)].width = w
wb.save(salida)
print(len(rows), 'filas')
