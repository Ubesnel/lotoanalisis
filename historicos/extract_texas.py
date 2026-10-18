"""Histórico de Texas (Pick 3 + Daily 4) a Excel, listo para el wizard.

Uso: python extract_texas.py <carpeta_cache> <salida.xlsx> [--refrescar]

Fuente oficial de Texas Lottery, un CSV por juego y turno (se descargan solos si no
están en la carpeta o con --refrescar):
  https://www.texaslottery.com/export/sites/lottery/Games/Pick_3/Winning_Numbers/pick3<turno>.csv
  https://www.texaslottery.com/export/sites/lottery/Games/Daily_4/Winning_Numbers/daily4<turno>.csv
con turno morning (10:00 → Mañana), day (12:27 → Mediodía), evening (18:00 → Tarde)
y night (22:12 → Noche). Cada fila: juego, mes, día, año, dígitos (3 o 4), suma, Fireball.
La Fireball se ignora. Pick 3 da la centena y el número (últimos dos dígitos); Daily 4
da los premios 2 y 3 (primer y segundo par). Como el wizard exige los premios 2 y 3, solo
salen las fechas con los dos juegos: Mediodía y Noche desde el 01/10/2007 (arranque de
Daily 4) y Mañana y Tarde desde el 09/09/2013. Se sortea de lunes a sábado.
"""
import csv
import os
import sys
from collections import Counter
from datetime import date, timedelta

import requests
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

BASE = 'https://www.texaslottery.com/export/sites/lottery/Games/%s/Winning_Numbers/%s%s.csv'
JUEGOS = {3: ('Pick_3', 'pick3'), 4: ('Daily_4', 'daily4')}
TURNOS = {'morning': ('morning', 'Mañana'), 'day': ('midday', 'Mediodía'),
          'evening': ('afternoon', 'Tarde'), 'night': ('evening', 'Noche')}
ORDEN = [c for c, _ in TURNOS.values()]


def cargar(carpeta, refrescar):
    os.makedirs(carpeta, exist_ok=True)
    datos = {}   # (juego, turno) -> {fecha: 'dígitos'}
    for juego, (carp, pref) in JUEGOS.items():
        for turno in TURNOS:
            ruta = os.path.join(carpeta, '%s%s.csv' % (pref, turno))
            if refrescar or not os.path.exists(ruta):
                r = requests.get(BASE % (carp, pref, turno), timeout=60,
                                 headers={'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) Chrome/147.0'})
                r.raise_for_status()
                with open(ruta, 'wb') as f:
                    f.write(r.content)
            d = {}
            with open(ruta, encoding='utf-8-sig', newline='') as f:
                for fila in csv.reader(f):
                    if len(fila) < 4 + juego:
                        continue
                    dig = ''.join(x.strip() for x in fila[4:4 + juego])
                    if len(dig) == juego and dig.isdigit():
                        d[date(int(fila[3]), int(fila[1]), int(fila[2]))] = dig
            datos[(juego, turno)] = d
    return datos


def main():
    carpeta, salida = sys.argv[1], sys.argv[2]
    datos = cargar(carpeta, '--refrescar' in sys.argv)
    hoy = date.today()

    filas, obs = [], []
    for turno, (codigo, nombre) in TURNOS.items():
        p3, p4 = datos[(3, turno)], datos[(4, turno)]
        d = max(min(p3), min(p4))
        while d <= hoy:
            if d in p3 and d in p4:
                filas.append((d, codigo, nombre, p3[d], p4[d]))
            elif d < hoy and d.weekday() != 6:
                faltan = 'Pick 3 y Daily 4' if d not in p3 and d not in p4 else \
                    ('Pick 3' if d not in p3 else 'Daily 4')
                obs.append((d, codigo, 'sin resultado en la fuente (se omite el turno)', faltan))
            d += timedelta(days=1)

    filas.sort(key=lambda r: (r[0], ORDEN.index(r[1])))
    wb = Workbook()
    ws = wb.active
    ws.title = 'Texas'
    ws.append(['Fecha', 'Turno', 'Turno (nombre)', 'Pick 3', 'Daily 4', 'Centena', 'Número',
               'Premio 2', 'Premio 3', 'Fuente'])
    for fecha, codigo, nombre, p3, p4 in filas:
        ws.append([fecha, codigo, nombre, p3, p4, int(p3[0]), p3[1:], p4[:2], p4[2:], 'texaslottery.com'])
    for c in ws[1]:
        c.font = Font(bold=True)
    for row in ws.iter_rows(min_row=2, min_col=1, max_col=1):
        row[0].number_format = 'yyyy-mm-dd'
    for col in (4, 5, 7, 8, 9):
        for row in ws.iter_rows(min_row=2, min_col=col, max_col=col):
            row[0].number_format = '@'
    for i, w in enumerate([12, 10, 15, 8, 8, 9, 9, 9, 9, 16], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = 'A2'

    wo = wb.create_sheet('Observaciones')
    wo.append(['Fecha', 'Turno', 'Observación', 'Detalle'])
    for c in wo[1]:
        c.font = Font(bold=True)
    for o in sorted(obs, key=lambda x: (x[0], x[1])):
        wo.append(list(o))
    for row in wo.iter_rows(min_row=2, min_col=1, max_col=1):
        row[0].number_format = 'yyyy-mm-dd'
    for i, w in enumerate([12, 10, 46, 16], 1):
        wo.column_dimensions[get_column_letter(i)].width = w
    wb.save(salida)

    print('filas', len(filas), 'desde', filas[0][0], 'hasta', filas[-1][0])
    print('por turno', Counter(r[1] for r in filas))
    print('observaciones', len(obs), sorted(Counter(o[0].year for o in obs).items()))


if __name__ == '__main__':
    main()
