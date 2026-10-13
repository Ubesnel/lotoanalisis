"""Histórico de La Diaria de El Salvador a Excel, listo para el wizard.

Uso: python extract_la_diaria_el_salvador.py <cache.json> <salida.xlsx>

Fuente oficial: loto.sv (LOTO, Lotería Nacional de Beneficencia), el mismo
endpoint que usa el selector de fechas de su página de La Diaria:
  GET https://loto.sv/api/resultados_diaria_sv.php?fecha=YYYY-MM-DD
  → {"11": {"hora": "10:58:00…", "par1": "71"}, "18": {…} | null, "21": …}
Un día por petición. El caché JSON guarda lo ya bajado; los últimos días se
vuelven a pedir siempre, por si el sitio cargó tarde algún sorteo.

Horarios (hora de El Salvador), según los datos:
  · desde el 17/07/2024: 11:00 y 21:00 todos los días.
  · desde el 24/08/2026: además 18:00.
Antes del 17/07/2024 hay un único sorteo suelto (22/04/2024) que se descarta:
dejaría un hueco de tres meses al principio y distorsionaría los atrasos.
Un sorteo de 18:00 anterior al 24/08/2026 (pasó el 19/10/2025, el día que
falta el de 11:00) se toma como el de 11:00 cargado en el turno equivocado.
"""
import json
import os
import sys
import time
from collections import Counter
from datetime import date, timedelta

import requests
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

URL = 'https://loto.sv/api/resultados_diaria_sv.php'
DESDE = date(2024, 7, 17)
INICIO_1800 = date(2026, 8, 24)
REPEDIR_DIAS = 3
DELAY = 0.7

TURNO = {'11': 'h1100', '18': 'h1800', '21': 'h2100'}
NOMBRE = {'h1100': '11:00 AM', 'h1800': '6:00 PM', 'h2100': '9:00 PM'}
ORDEN = {t: i for i, t in enumerate(NOMBRE)}


def descargar(cache_path, hasta):
    cache = {}
    if os.path.exists(cache_path):
        with open(cache_path, encoding='utf-8') as f:
            cache = json.load(f)
    s = requests.Session()
    s.headers.update({'User-Agent': ('Mozilla/5.0 (Windows NT 10.0; Win64; x64) '
                                     'AppleWebKit/537.36 (KHTML, like Gecko) Chrome/147.0 Safari/537.36'),
                      'Referer': 'https://loto.sv/index.php?pag=diaria'})
    repedir = hasta - timedelta(days=REPEDIR_DIAS)
    d = DESDE
    while d <= hasta:
        k = d.isoformat()
        if k not in cache or 'error' in cache[k] or d > repedir:
            for intento in range(3):
                try:
                    r = s.get(URL, params={'fecha': k}, timeout=40)
                    r.raise_for_status()
                    cache[k] = r.json()
                    break
                except Exception as exc:
                    cache[k] = {'error': str(exc)}
                    time.sleep(5 * (intento + 1))
            time.sleep(DELAY)
        d += timedelta(days=1)
    with open(cache_path, 'w', encoding='utf-8') as f:
        json.dump(cache, f)
    return cache


def main():
    cache_path, salida = sys.argv[1], sys.argv[2]
    hasta = date.today() - timedelta(days=1)
    cache = descargar(cache_path, hasta)

    data, obs = {}, []
    d = DESDE
    while d <= hasta:
        dia = cache.get(d.isoformat()) or {}
        if 'error' in dia:
            obs.append((d, '', 'error de descarga', dia['error']))
        for clave, turno in TURNO.items():
            v = dia.get(clave)
            if not v:
                continue
            par = str(v.get('par1', '')).strip()
            if not par.isdigit() or int(par) > 99:
                obs.append((d, turno, 'número ilegible', par))
                continue
            if turno == 'h1800' and d < INICIO_1800:
                if (d, 'h1100') in data or (dia.get('11') or None):
                    obs.append((d, turno, 'sorteo de 18:00 antes de que exista (descartado)', par.zfill(2)))
                    continue
                obs.append((d, 'h1100', 'cargado como 18:00, se toma como el de 11:00', par.zfill(2)))
                turno = 'h1100'
            data[(d, turno)] = {'fecha': d, 'turno': turno,
                                'hora': str(v.get('hora', ''))[:5], 'numero': int(par)}
        esperados = ['h1100', 'h2100'] + (['h1800'] if d >= INICIO_1800 else [])
        for t in esperados:
            if (d, t) not in data:
                obs.append((d, t, 'turno faltante', ''))
        d += timedelta(days=1)

    filas = sorted(data.values(), key=lambda r: (r['fecha'], ORDEN[r['turno']]))
    wb = Workbook()
    ws = wb.active
    ws.title = 'La Diaria'
    ws.append(['Fecha', 'Turno', 'Turno (nombre)', 'Hora', 'Número', 'Fuente'])
    for r in filas:
        ws.append([r['fecha'], r['turno'], NOMBRE[r['turno']], r['hora'], '%02d' % r['numero'], 'loto.sv'])
    for c in ws[1]:
        c.font = Font(bold=True)
    for row in ws.iter_rows(min_row=2, min_col=1, max_col=1):
        row[0].number_format = 'yyyy-mm-dd'
    for row in ws.iter_rows(min_row=2, min_col=5, max_col=5):
        row[0].number_format = '@'
    for i, w in enumerate([12, 9, 15, 8, 9, 10], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = 'A2'

    wo = wb.create_sheet('Observaciones')
    wo.append(['Fecha', 'Turno', 'Observación', 'Detalle'])
    for c in wo[1]:
        c.font = Font(bold=True)
    for o in sorted(obs, key=lambda x: (x[0], ORDEN.get(x[1], 9))):
        wo.append(list(o))
    for row in wo.iter_rows(min_row=2, min_col=1, max_col=1):
        row[0].number_format = 'yyyy-mm-dd'
    for i, w in enumerate([12, 9, 46, 12], 1):
        wo.column_dimensions[get_column_letter(i)].width = w
    wb.save(salida)

    print('filas', len(filas), 'desde', filas[0]['fecha'], 'hasta', filas[-1]['fecha'])
    print('por turno', Counter(r['turno'] for r in filas))
    print('observaciones', Counter(o[2] for o in obs))


if __name__ == '__main__':
    main()
