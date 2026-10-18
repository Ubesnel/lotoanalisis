"""Loteka Tarde (13:30, turno Tarde): arrancó el 28/09/2026, solo está en
enloteria. Cada página /resultados-loteka-tarde-<fecha> lista los 14 sorteos
anteriores a esa fecha. Uso: python extract_loteka_tarde.py <salida.xlsx>"""
import html
import re
import sys
import time
from datetime import date, timedelta

import requests
from openpyxl import Workbook

MESES = {m: i + 1 for i, m in enumerate(
    'enero febrero marzo abril mayo junio julio agosto septiembre octubre noviembre diciembre'.split())}
PAT = re.compile(r'Loteka Tarde \w+ (\d{1,2}) de (\w+), (\d{4}) \d{1,2}:\d{2}\s*[AP]M (\d{2}) (\d{2}) (\d{2})')
INICIO = date(2026, 9, 28)

out, d = {}, date.today()
while d >= INICIO:
    r = requests.get('https://enloteria.com/resultados-loteka-tarde-%s' % d, timeout=40,
                     headers={'User-Agent': 'Mozilla/5.0'})
    time.sleep(1.5)
    if r.status_code != 200:
        d -= timedelta(days=1)
        continue
    t = re.sub(r'(?s)<(script|style).*?</\1>', '', r.text)
    t = re.sub(r'\s+', ' ', html.unescape(re.sub(r'<[^>]+>', ' ', t)))
    for m in PAT.finditer(t):
        f = date(int(m.group(3)), MESES[m.group(2).lower()], int(m.group(1)))
        if f >= INICIO:
            out[f] = m.groups()[3:]
    d = min(out) - timedelta(days=1) if out and min(out) <= d else d - timedelta(days=1)
wb = Workbook()
ws = wb.active
ws.title = 'Salidas'
ws.append(['Fecha', 'Turno', 'Turno (nombre)', 'Número', 'Premio 2', 'Premio 3', 'Fuente'])
for f in sorted(out):
    ws.append([f, 'afternoon', 'Tarde', *out[f], 'enloteria'])
for row in ws.iter_rows(min_row=2, max_col=1):
    row[0].number_format = 'yyyy-mm-dd'
for row in ws.iter_rows(min_row=2, min_col=4, max_col=6):
    for c in row:
        c.number_format = '@'
wb.save(sys.argv[1])
print(len(out), min(out), max(out))
