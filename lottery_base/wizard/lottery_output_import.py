# -*- coding: utf-8 -*-
import base64
import io
import logging
from datetime import date, datetime

from odoo import fields, models
from odoo.exceptions import UserError

_logger = logging.getLogger(__name__)

# Encabezados que se buscan en la primera fila (sin importar mayúsculas ni
# tildes). Las demás columnas del Excel se ignoran.
COL_FECHA = 'fecha'
COL_TURNO = 'turno'
COL_NUMERO = 'numero'
COL_CENTENA = 'centena'
COL_PREMIO_2 = 'premio 2'
COL_PREMIO_3 = 'premio 3'

_MAX_DETAIL_LINES = 100


def _norm(texto):
    texto = str(texto or '').strip().lower()
    for a, b in (('á', 'a'), ('é', 'e'), ('í', 'i'), ('ó', 'o'), ('ú', 'u')):
        texto = texto.replace(a, b)
    return texto


class LotteryOutputImport(models.TransientModel):
    _name = 'lottery.output.import'
    _description = 'Importar salidas desde Excel'

    sorteo_id = fields.Many2one('lottery.sorteo', string='Sorteo', required=True)
    file = fields.Binary(string='Archivo Excel (.xlsx)', required=True)
    filename = fields.Char(string='Nombre del archivo')
    complete_premios = fields.Boolean(
        string='Completar premios faltantes',
        help='En las salidas que ya existen y tienen vacío el premio 2 o el premio 3, '
             'carga el que trae el Excel. No toca el número, la centena ni los premios ya cargados.')
    # El resultado va en campos simples y la vista arma el aviso: un campo
    # Html obliga a cargar el editor de Odoo (web_editor) solo para mostrarlo.
    state = fields.Selection([
        ('draft', 'Sin correr'),
        ('validated', 'Validado'),
        ('imported', 'Importado'),
        ('error', 'Con errores'),
    ], default='draft', readonly=True)
    summary = fields.Char(string='Resumen', readonly=True)
    errors = fields.Text(string='Errores', readonly=True)

    def action_validate(self):
        self.ensure_one()
        self.write(self._run(dry_run=True))
        return self._reopen()

    def action_import(self):
        self.ensure_one()
        self.write(self._run(dry_run=False))
        return self._reopen()

    def _reopen(self):
        return {
            'type': 'ir.actions.act_window',
            'res_model': self._name,
            'res_id': self.id,
            'view_mode': 'form',
            'target': 'new',
        }

    # ── Lectura ───────────────────────────────────────────────────

    def _read_rows(self):
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise UserError('Falta la librería openpyxl.') from exc
        try:
            wb = load_workbook(io.BytesIO(base64.b64decode(self.file)),
                               read_only=True, data_only=True)
        except Exception as exc:
            raise UserError(f'No se pudo leer el archivo como Excel (.xlsx): {exc}') from exc

        # Primera hoja: es la que trae los datos en el Excel de extracción.
        filas = wb.worksheets[0].iter_rows(values_only=True)
        encabezado = [_norm(c) for c in next(filas, ())]
        faltan = [c for c in (COL_FECHA, COL_TURNO, COL_NUMERO) if c not in encabezado]
        if self.sorteo_id.uses_hundreds and COL_CENTENA not in encabezado:
            faltan.append(COL_CENTENA)
        # Quinielas (premios 2 y 3): los corridos son parte del resultado,
        # así que si el sorteo los usa el Excel tiene que traerlos.
        if self.sorteo_id.is_pick3:
            faltan += [c for c in (COL_PREMIO_2, COL_PREMIO_3) if c not in encabezado]
        if faltan:
            raise UserError('Faltan columnas en la primera fila: %s.\nEncontradas: %s'
                            % (', '.join(faltan), ', '.join(c for c in encabezado if c)))
        idx = {c: encabezado.index(c)
               for c in (COL_FECHA, COL_TURNO, COL_NUMERO, COL_CENTENA, COL_PREMIO_2, COL_PREMIO_3)
               if c in encabezado}

        for nro, fila in enumerate(filas, start=2):
            if not fila or all(v in (None, '') for v in fila):
                continue
            yield nro, {c: (fila[i] if i < len(fila) else None) for c, i in idx.items()}

    @staticmethod
    def _parse_date(valor):
        if isinstance(valor, datetime):
            return valor.date()
        if isinstance(valor, date):
            return valor
        texto = str(valor or '').strip()[:10]
        for fmt in ('%Y-%m-%d', '%d/%m/%Y'):
            try:
                return datetime.strptime(texto, fmt).date()
            except ValueError:
                continue
        return None

    @staticmethod
    def _parse_int(valor):
        try:
            return int(str(valor).strip().split('.')[0])
        except (TypeError, ValueError):
            return None

    # ── Importación ───────────────────────────────────────────────

    def _run(self, dry_run):
        sorteo = self.sorteo_id
        # Se acepta el código del turno (morning) o su nombre (Mañana), pero
        # solo turnos que el sorteo juega.
        turnos = {}
        for t in sorteo.turno_ids:
            turnos[_norm(t.code)] = t
            turnos[_norm(t.name)] = t
        numeros = {n['name']: n['id']
                   for n in self.env['lottery.number'].search_read([], ['name'])}

        Output = self.env['lottery.output']
        existentes = {
            (fields.Date.to_date(r['date']), r['turno_id'][0]): r
            for r in Output.search_read([('sorteo_id', '=', sorteo.id)],
                                        ['date', 'turno_id', 'premio_2_id', 'premio_3_id'])
        }
        completar = self.complete_premios and sorteo.is_pick3

        vals_list, errores, omitidas = [], [], 0
        completadas = []   # (id, vals) de salidas existentes a las que les faltan premios
        vistos = set()
        for nro, fila in self._read_rows():
            fecha = self._parse_date(fila[COL_FECHA])
            turno = turnos.get(_norm(fila[COL_TURNO]))
            numero = self._parse_int(fila[COL_NUMERO])
            if not fecha:
                errores.append(f'Fila {nro}: fecha inválida ({fila[COL_FECHA]}).')
                continue
            if not turno:
                errores.append(f'Fila {nro}: turno "{fila[COL_TURNO]}" no es de {sorteo.name}.')
                continue
            if numero not in numeros:
                errores.append(f'Fila {nro}: número inválido ({fila[COL_NUMERO]}).')
                continue
            clave = (fecha, turno.id)
            if clave in vistos:
                errores.append(f'Fila {nro}: {fecha} {turno.name} repetido en el archivo.')
                continue
            vistos.add(clave)
            if clave in existentes:
                if completar:
                    actual = existentes[clave]
                    faltan = {}
                    for campo, col in (('premio_2_id', COL_PREMIO_2), ('premio_3_id', COL_PREMIO_3)):
                        if actual[campo]:
                            continue
                        p = self._parse_int(fila.get(col))
                        if p not in numeros:
                            errores.append(f'Fila {nro}: {col} inválido ({fila.get(col)}).')
                            break
                        faltan[campo] = numeros[p]
                    else:
                        if faltan:
                            completadas.append((actual['id'], faltan))
                            continue
                omitidas += 1
                continue
            vals = {'sorteo_id': sorteo.id, 'date': fecha,
                    'turno_id': turno.id, 'number_id': numeros[numero]}
            if sorteo.uses_hundreds:
                centena = self._parse_int(fila.get(COL_CENTENA))
                if centena not in numeros:
                    errores.append(f'Fila {nro}: centena inválida ({fila.get(COL_CENTENA)}).')
                    continue
                vals['hundreds_id'] = numeros[centena]
            if sorteo.is_pick3:
                p2 = self._parse_int(fila.get(COL_PREMIO_2))
                p3 = self._parse_int(fila.get(COL_PREMIO_3))
                if p2 not in numeros or p3 not in numeros:
                    errores.append(f'Fila {nro}: premios 2 y 3 inválidos '
                                   f'({fila.get(COL_PREMIO_2)}, {fila.get(COL_PREMIO_3)}).')
                    continue
                vals['premio_2_id'] = numeros[p2]
                vals['premio_3_id'] = numeros[p3]
            vals_list.append(vals)

        # Con errores no se importa nada: se corrige el Excel y se vuelve a
        # subir, así no queda un histórico a medias.
        if errores or dry_run:
            return self._resultado(sorteo, len(vals_list), omitidas, errores, importado=False,
                                   completadas=len(completadas))

        # Un write por (campo, número) y no por salida: son miles de filas y
        # uno por uno excede el tiempo máximo de la petición.
        por_valor = {}
        for out_id, vals in completadas:
            for campo, numero_id in vals.items():
                por_valor.setdefault((campo, numero_id), []).append(out_id)
        for (campo, numero_id), ids in por_valor.items():
            Output.browse(ids).with_context(skip_prediction_validation=True).write({campo: numero_id})
        if completadas:
            _logger.info('Importación Excel (%s): %d salidas con premios completados.',
                         sorteo.name, len(completadas))

        if vals_list:
            # Mismos contextos que los backfills de los importadores: no
            # recalcular el próximo sorteo por cada salida ni validar
            # predicciones con salidas viejas.
            Output.with_context(skip_next_draw_recompute=True,
                                skip_prediction_validation=True).create(vals_list)
            sorteo._recompute_next_draw()
            _logger.info('Importación Excel (%s): %d salidas creadas.', sorteo.name, len(vals_list))
        return self._resultado(sorteo, len(vals_list), omitidas, errores, importado=True,
                               completadas=len(completadas))

    @staticmethod
    def _resultado(sorteo, nuevas, omitidas, errores, importado, completadas=0):
        extra = f' Con premios completados: {completadas}.' if completadas else ''
        if errores:
            state, titulo = 'error', f'Hay {len(errores)} errores: no se importó nada.'
        elif importado:
            state, titulo = 'imported', f'Importadas {nuevas} salidas de {sorteo.name}.{extra}'
        else:
            state, titulo = 'validated', (f'Validación OK: se importarían {nuevas} salidas de '
                                          f'{sorteo.name}.{extra}')
        detalle = errores[:_MAX_DETAIL_LINES]
        if len(errores) > _MAX_DETAIL_LINES:
            detalle.append(f'… y {len(errores) - _MAX_DETAIL_LINES} más.')
        return {
            'state': state,
            'summary': f'{titulo} Ya registradas (se saltean): {omitidas}.',
            'errors': '\n'.join(detalle) or False,
        }
